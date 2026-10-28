import uuid
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook
from pypdf import PdfReader
from django.test import TestCase, TransactionTestCase, override_settings, Client
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission, Group
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.utils import timezone
from .models import (Worker, PayrollSettings, Attendance, Payroll, Audit, SprayProduct,
                     SprayApplication, Customer, Product, PaymentMethod, Sale)
from .farm_forms import WorkerForm, SettingsForm, SprayProductForm, SprayForm, SprayRows, AttendanceRows
from .farm_services import (save_catalog, save_attendance, generate_payroll, payroll_summary,
                            payroll_transition, save_spray, cancel_spray)
from .farm_exports import payroll_xlsx
from .services import create_sale
from .receipt import whatsapp_url


class FarmTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser('farm-admin', password='test-password')
        cls.viewer = get_user_model().objects.create_user('viewer', password='test-password')
        cls.day = timezone.localdate() - timedelta(days=10)
        cls.a = Worker.objects.create(name='Ana', joined_on=cls.day, daily_rate=50000)
        cls.b = Worker.objects.create(name='Beatriz', joined_on=cls.day, daily_rate=60000)
        PayrollSettings.objects.create(pk=1, mode='hourly', hourly_rate=10000)
        cls.p1 = SprayProduct.objects.create(name='Producto ensayo A', unit='ml', interval_days=7, interval_source='Responsable de ensayo')
        cls.p2 = SprayProduct.objects.create(name='Producto ensayo B', unit='g')

    def setUp(self):
        self.client.force_login(self.admin)

    def row(self, worker, worked=True, hours=0, minutes=0, day=None, notes=''):
        record = Attendance.objects.filter(worker=worker, date=day or self.day).first()
        return dict(worker=worker.pk, worked=worked, hours=hours, minutes=minutes, notes=notes,
                    revision=record.revision if record else 0)

    def generate(self, start=None, end=None, token=None):
        return generate_payroll(start or self.day, end or self.day, token or uuid.uuid4(), self.admin)

    def test_different_rates_absent_and_ninety_minutes(self):
        save_attendance(self.day, [self.row(self.a, minutes=90), self.row(self.b, worked=False)], self.admin)
        p = self.generate()
        self.assertEqual(p.total, Decimal('65000.00'))
        summary = payroll_summary(p)
        self.assertEqual(summary[0]['hours'], 1)
        self.assertEqual(summary[0]['remaining_minutes'], 30)
        self.assertEqual(summary[1]['days'], 0)
        self.assertEqual(summary[1]['total'], 0)

    def test_rates_preserved_on_edit_and_after_new_rates(self):
        save_attendance(self.day, [self.row(self.a, hours=1)], self.admin)
        self.a.daily_rate = 90000
        self.a.save()
        PayrollSettings.objects.filter(pk=1).update(hourly_rate=20000)
        save_attendance(self.day, [self.row(self.a, hours=2)], self.admin)
        day2 = self.day + timedelta(days=1)
        save_attendance(day2, [self.row(self.a, hours=1, day=day2)], self.admin)
        p = self.generate(end=day2)
        self.assertEqual(p.total, Decimal('180000.00'))
        self.assertEqual(p.days.first().snapshot['daily_rate'], '50000.00')

    def test_derived_tariff_rounds_only_final_amount(self):
        PayrollSettings.objects.filter(pk=1).update(mode='derived', ordinary_hours=3, multiplier=1)
        self.a.daily_rate = Decimal('.01')
        self.a.save()
        save_attendance(self.day, [self.row(self.a, minutes=90)], self.admin)
        p = self.generate()
        self.assertEqual(p.days.get().snapshot['extra'], '0.01')

    def test_partial_ranges_and_no_overlap(self):
        for offset in (0, 1, 2):
            day = self.day + timedelta(days=offset)
            save_attendance(day, [self.row(self.b, day=day)], self.admin)
        first = self.generate(end=self.day + timedelta(days=1))
        self.assertEqual(first.total, 120000)
        with self.assertRaises(ValidationError):
            self.generate(start=self.day + timedelta(days=1), end=self.day + timedelta(days=2))
        self.assertEqual(self.generate(start=self.day + timedelta(days=2), end=self.day + timedelta(days=2)).total, 60000)

    def test_duplicate_attendance_stale_edit_and_database_unique(self):
        row = self.row(self.a)
        save_attendance(self.day, [row], self.admin)
        with self.assertRaises(ValidationError):
            save_attendance(self.day, [row], self.admin)
        with self.assertRaises(ValidationError):
            save_attendance(self.day, [self.row(self.a), self.row(self.a)], self.admin)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Attendance.objects.create(worker=self.a, date=self.day, worked=True, daily_rate=1, hourly_rate=0, created_by=self.admin)
        self.assertEqual(Attendance.objects.count(), 1)

    def test_future_before_join_and_absent_extras_rejected(self):
        for day, worked, minutes in [(timezone.localdate()+timedelta(days=1), True, 0), (self.day-timedelta(days=1), True, 0), (self.day, False, 1)]:
            with self.subTest(day=day, worked=worked), self.assertRaises(ValidationError):
                save_attendance(day, [self.row(self.a, worked=worked, minutes=minutes, day=day)], self.admin)
        self.assertFalse(Attendance.objects.exists())

    def test_inactive_worker_and_missing_configuration(self):
        self.a.active = False
        self.a.save()
        with self.assertRaises(ValidationError):
            save_attendance(self.day, [self.row(self.a)], self.admin)
        PayrollSettings.objects.all().delete()
        with self.assertRaises(ValidationError):
            save_attendance(self.day, [self.row(self.b)], self.admin)

    def test_payroll_retry_payment_lock_and_controlled_correction(self):
        save_attendance(self.day, [self.row(self.a)], self.admin)
        token = uuid.uuid4()
        p = self.generate(token=token)
        self.assertEqual(self.generate(token=token).pk, p.pk)
        payroll_transition(p.pk, 'pay', '', self.admin)
        original = list(p.days.values_list('snapshot', flat=True))
        with self.assertRaises(ValidationError):
            save_attendance(self.day, [self.row(self.a, hours=1)], self.admin)
        with self.assertRaises(ValidationError):
            payroll_transition(p.pk, 'cancel', 'Corrección', self.viewer)
        with self.assertRaises(ValidationError):
            payroll_transition(p.pk, 'cancel', '', self.admin)
        payroll_transition(p.pk, 'cancel', 'Corrección de horas', self.admin)
        save_attendance(self.day, [self.row(self.a, hours=1)], self.admin)
        self.assertEqual(list(p.days.values_list('snapshot', flat=True)), original)
        self.assertEqual(self.generate().total, 60000)
        self.assertTrue(Audit.objects.filter(action='Estado de nómina').exists())

    def test_locked_unchanged_row_allows_other_workers(self):
        save_attendance(self.day, [self.row(self.a)], self.admin)
        self.generate()
        save_attendance(self.day, [self.row(self.a), self.row(self.b)], self.admin)
        self.assertEqual(Attendance.objects.count(), 2)

    def test_invalid_empty_payroll_and_period_form(self):
        with self.assertRaises(ValidationError):
            self.generate()
        response = self.client.post('/nomina/generar/', {'start': self.day, 'end': self.day-timedelta(days=1), 'token': uuid.uuid4()})
        self.assertContains(response, 'La fecha final')
        self.assertFalse(Payroll.objects.exists())

    def test_xlsx_two_sheets_reconcile_and_text_injection(self):
        self.a.name = '=SUM(A1:A2)'
        self.a.save()
        save_attendance(self.day, [self.row(self.a, minutes=90), self.row(self.b)], self.admin)
        p = self.generate()
        response = payroll_xlsx(p)
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ['Resumen', 'Detalle diario'])
        self.assertEqual(workbook['Resumen']['A3'].data_type, 's')
        self.assertEqual(Decimal(str(workbook['Resumen']['F5'].value)), p.total)
        self.assertEqual(sum(Decimal(str(r[9])) for r in workbook['Detalle diario'].iter_rows(min_row=2, values_only=True)), p.total)
        self.assertEqual(workbook['Detalle diario']['E2'].value, 90)

    def spray_data(self, **overrides):
        data = dict(scheduled_at=(timezone.localtime()-timedelta(days=2)).strftime('%Y-%m-%dT%H:%M'),
                    area='Lote norte', responsible='Responsable de prueba', status='completed', notes='',
                    token=str(uuid.uuid4()), revision=0, next_at='')
        data.update(overrides)
        return data

    def spray_rows(self):
        return [dict(product=self.p1, quantity=Decimal('12.250'), unit='ml'),
                dict(product=self.p2, quantity=Decimal('3'), unit='g')]

    def spray_form(self, data=None, instance=None):
        form = SprayForm(data or self.spray_data(), instance=instance)
        self.assertTrue(form.is_valid(), form.errors)
        return form

    def test_spray_multiple_products_and_future_schedule_no_auto_completion(self):
        next_at = timezone.localtime()+timedelta(days=3)
        form = self.spray_form(self.spray_data(next_at=next_at.strftime('%Y-%m-%dT%H:%M')))
        obj = save_spray(form, self.spray_rows(), self.admin)
        self.assertEqual(obj.lines.count(), 2)
        followup = obj.followups.get()
        self.assertEqual(followup.status, 'scheduled')
        self.assertIsNone(followup.completed_at)
        SprayApplication.objects.filter(pk=followup.pk).update(scheduled_at=timezone.now()-timedelta(days=1))
        followup.refresh_from_db()
        self.assertTrue(followup.overdue)
        self.assertEqual(followup.status, 'scheduled')
        self.assertEqual(save_spray(self.spray_form(dict(form.data)), self.spray_rows(), self.admin).pk, obj.pk)
        self.assertEqual(SprayApplication.objects.count(), 2)

    def test_spray_filters_calendar_and_history(self):
        obj = save_spray(self.spray_form(), self.spray_rows(), self.admin)
        for path in ['/fumigacion/', '/fumigacion/calendario/']:
            response = self.client.get(path, {'area': 'norte', 'product': self.p1.pk})
            self.assertContains(response, 'Producto ensayo A')
            self.assertContains(response, 'Lote norte')
            self.assertNotContains(self.client.get(path, {'area': 'otro'}), 'Lote norte')
        self.assertContains(self.client.get('/fumigacion/', {'start': timezone.localdate().isoformat()}), 'No hay aplicaciones')
        self.assertEqual(self.client.get(f'/fumigacion/{obj.pk}/').status_code, 200)

    def test_confirm_schedule_explicit_and_completed_immutable(self):
        obj = save_spray(self.spray_form(self.spray_data(status='scheduled')), self.spray_rows(), self.admin)
        form = self.spray_form(self.spray_data(token=obj.token, revision=obj.revision), instance=obj)
        obj = save_spray(form, self.spray_rows(), self.admin)
        self.assertEqual(obj.status, 'completed')
        self.assertIsNotNone(obj.completed_at)
        with self.assertRaises(ValidationError):
            save_spray(form, self.spray_rows(), self.admin)
        self.assertRedirects(self.client.get(f'/fumigacion/{obj.pk}/editar/'), f'/fumigacion/{obj.pk}/')

    def test_spray_full_http_post_with_two_products(self):
        data = self.spray_data()
        data.update({'form-TOTAL_FORMS': 2, 'form-INITIAL_FORMS': 0,
                     'form-0-product': self.p1.pk, 'form-0-quantity': '12.25', 'form-0-unit': 'ml',
                     'form-1-product': self.p2.pk, 'form-1-quantity': '3', 'form-1-unit': 'g'})
        response = self.client.post('/fumigacion/nueva/', data)
        self.assertEqual(response.status_code, 302)
        obj = SprayApplication.objects.get()
        self.assertEqual(obj.lines.count(), 2)
        self.assertEqual(self.client.post('/fumigacion/nueva/', data).status_code, 302)
        self.assertEqual(SprayApplication.objects.count(), 1)

    def test_cancel_scheduled_requires_reason_and_audit(self):
        obj = save_spray(self.spray_form(self.spray_data(status='scheduled')), self.spray_rows(), self.admin)
        with self.assertRaises(ValidationError):
            cancel_spray(obj.pk, '', self.admin)
        cancel_spray(obj.pk, 'Cambio de condiciones', self.admin)
        obj.refresh_from_db()
        self.assertEqual(obj.status, 'cancelled')
        self.assertEqual(obj.lines.count(), 2)
        self.assertTrue(Audit.objects.filter(action='Anular programación').exists())

    def test_duplicate_spray_products_and_invalid_forms(self):
        with self.assertRaises(ValidationError):
            save_spray(self.spray_form(), [self.spray_rows()[0]]*2, self.admin)
        self.assertFalse(SprayApplication.objects.exists())
        form = SprayForm(self.spray_data(scheduled_at=(timezone.localtime()+timedelta(days=1)).strftime('%Y-%m-%dT%H:%M')))
        self.assertFalse(form.is_valid())
        form = SprayProductForm({'name': 'Test', 'unit': 'ml', 'interval_days': 7, 'interval_source': '', 'revision': 0})
        self.assertFalse(form.is_valid())
        rows = SprayRows({'form-TOTAL_FORMS': 1, 'form-INITIAL_FORMS': 0,
                         'form-0-product': self.p1.pk, 'form-0-quantity': '1.25', 'form-0-unit': 'ml'})
        self.assertTrue(rows.is_valid(), rows.errors)

    def test_catalog_revision_settings_no_assumed_rates(self):
        form = WorkerForm(dict(name='Ana editada', joined_on=self.day, daily_rate=70000, active=True, revision=0), instance=Worker.objects.get(pk=self.a.pk))
        self.assertTrue(form.is_valid(), form.errors)
        save_catalog(form, self.admin)
        with self.assertRaises(ValidationError):
            save_catalog(form, self.admin)
        invalid = SettingsForm({'mode': 'hourly', 'revision': 0})
        self.assertFalse(invalid.is_valid())
        invalid = SettingsForm({'mode': 'derived', 'revision': 0, 'ordinary_hours': 8})
        self.assertFalse(invalid.is_valid())
        PayrollSettings.objects.all().delete()
        response = self.client.post('/nomina/tarifas/', {'mode': 'hourly', 'hourly_rate': 0, 'revision': 0})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(PayrollSettings.objects.get(pk=1).hourly_rate, 0)

    def test_all_new_pages_render_and_bad_filters_safe(self):
        for url in ['/nomina/', '/nomina/trabajadores/', '/nomina/trabajadores/nuevo/',
                    f'/nomina/trabajadores/{self.a.pk}/', f'/nomina/trabajadores/{self.a.pk}/editar/',
                    '/nomina/asistencia/', '/nomina/generar/', '/nomina/tarifas/', '/fumigacion/',
                    '/fumigacion/productos/', '/fumigacion/productos/nuevo/', '/fumigacion/calendario/',
                    '/fumigacion/proximas/', '/fumigacion/nueva/']:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)
        for url in ['/nomina/asistencia/?date=bad', '/fumigacion/calendario/?date=bad', '/fumigacion/?start=bad', '/nomina/?end=bad']:
            self.assertEqual(self.client.get(url).status_code, 200)

    def test_attendance_post_tamper_and_valid_batch(self):
        data = {'form-TOTAL_FORMS': 2, 'form-INITIAL_FORMS': 2}
        for i, w in enumerate((self.a, self.b)):
            data.update({f'form-{i}-worker': w.pk, f'form-{i}-revision': 0, f'form-{i}-worked': 'on',
                         f'form-{i}-hours': 0, f'form-{i}-minutes': 90 if i == 0 else 0})
        response = self.client.post(f'/nomina/asistencia/?date={self.day}', data)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Attendance.objects.count(), 2)
        data['form-0-worker'] = 999999
        self.assertContains(self.client.post(f'/nomina/asistencia/?date={self.day}', data), 'Cambió el listado')

    def test_access_csrf_and_offline_guard(self):
        paths = ['/nomina/', '/nomina/trabajadores/', '/nomina/asistencia/', '/nomina/tarifas/', '/fumigacion/', '/fumigacion/nueva/']
        self.client.logout()
        for path in paths:
            self.assertEqual(self.client.get(path).status_code, 302)
        self.client.force_login(self.viewer)
        for path in paths:
            self.assertEqual(self.client.get(path).status_code, 403)
        secure = Client(enforce_csrf_checks=True)
        secure.force_login(self.admin)
        self.assertEqual(secure.post('/nomina/generar/', {}).status_code, 403)
        self.client.force_login(self.admin)
        with override_settings(OFFLINE_NODE=True):
            self.assertEqual(self.client.post('/nomina/asistencia/', {}).status_code, 403)
            self.assertEqual(self.client.get('/fumigacion/').status_code, 403)

    def test_operator_cannot_configure_rates(self):
        call_command('initialize', verbosity=0)
        self.viewer.groups.add(Group.objects.get(name='Operadores'))
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get('/nomina/').status_code, 200)
        self.assertEqual(self.client.get('/nomina/tarifas/').status_code, 403)


class ReceiptTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('receipt-admin', password='test-password')
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(name='Cliente comprobante', identification='ID-123', phone='3001234567')
        self.product = Product.objects.create(name='Canasta comprobante', unit='canasta', price=2000, by_weight=True, weight=16)
        self.cash = PaymentMethod.objects.create(name='Efectivo')
        self.credit = PaymentMethod.objects.create(name='Crédito', is_credit=True)

    def sale(self, credit):
        return create_sale(user=self.user, customer_id=self.customer.pk, rows=[dict(product=self.product.pk, quantity=2)],
            method=self.credit if credit else self.cash, initial='10000' if credit else '64000',
            cash_method=self.cash, token=uuid.uuid4())

    @override_settings(BUSINESS_NAME='Fresas de prueba')
    def test_cash_and_credit_pdf_saved_values_and_private_access(self):
        for credit in (False, True):
            sale = self.sale(credit)
            response = self.client.get(f'/ventas/{sale.pk}/?export=pdf')
            text = ' '.join(p.extract_text() for p in PdfReader(BytesIO(response.content)).pages)
            for label in ['Fresas de prueba', str(sale), 'Cliente comprobante', 'ID-123', 'Canasta comprobante', '$ 64.000,00']:
                self.assertIn(label, text)
            self.assertIn('$ 54.000,00' if credit else 'Saldo actual: $ 0,00', text)
            self.assertIn('private, no-store', response['Cache-Control'])
            self.assertIn('inline', self.client.get(f'/ventas/{sale.pk}/?export=pdf&view=1')['Content-Disposition'])
            detail = self.client.get(f'/ventas/{sale.pk}/')
            self.assertContains(detail, 'Enviar por WhatsApp')
            self.assertContains(detail, 'Abrir conversación de WhatsApp')
            self.assertIn('https://wa.me/573001234567', whatsapp_url(sale))
        self.assertEqual(Sale.objects.count(), 2)
        self.client.logout()
        self.assertEqual(self.client.get(f'/ventas/{sale.pk}/?export=pdf').status_code, 302)
        viewer = get_user_model().objects.create_user('no-receipt')
        self.client.force_login(viewer)
        self.assertEqual(self.client.get(f'/ventas/{sale.pk}/?export=pdf').status_code, 403)

    def test_invalid_phone_no_conversation_and_international(self):
        sale = self.sale(True)
        for phone in ['', '123', '3001234567 ext 5', '+000000000000']:
            sale.customer.phone = phone
            self.assertEqual(whatsapp_url(sale), '')
        sale.customer.phone = '+1 (202) 555-0100'
        self.assertIn('/12025550100?', whatsapp_url(sale))


class FarmConcurrencyTests(TransactionTestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('concurrent-admin')
        self.day = timezone.localdate()
        self.worker = Worker.objects.create(name='Concurrencia', joined_on=self.day, daily_rate=50000)
        PayrollSettings.objects.create(pk=1, mode='hourly', hourly_rate=10000)
        save_attendance(self.day, [dict(worker=self.worker.pk, worked=True, hours=0, minutes=0, notes='', revision=0)], self.user)

    def parallel(self, action):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        from django.db import close_old_connections
        barrier = Barrier(2)
        def run(i):
            close_old_connections()
            try:
                user = get_user_model().objects.get(pk=self.user.pk)
                barrier.wait(timeout=10)
                return action(i, user)
            except ValidationError:
                return 'rejected'
            finally:
                close_old_connections()
        with ThreadPoolExecutor(max_workers=2) as executor:
            return list(executor.map(run, (0, 1)))

    def test_parallel_overlapping_payrolls_only_one_wins(self):
        results = self.parallel(lambda i, user: generate_payroll(self.day, self.day, uuid.uuid4(), user).pk)
        self.assertEqual(results.count('rejected'), 1)
        self.assertEqual(Payroll.objects.count(), 1)

    def test_parallel_same_token_is_idempotent(self):
        token = uuid.uuid4()
        results = self.parallel(lambda i, user: generate_payroll(self.day, self.day, token, user).pk)
        self.assertEqual(results[0], results[1])
        self.assertEqual(Payroll.objects.count(), 1)

    def test_parallel_attendance_changes_reject_stale_revision(self):
        def action(i, user):
            save_attendance(self.day, [dict(worker=self.worker.pk, worked=True, hours=i+1, minutes=0, notes='', revision=1)], user)
            return 'saved'
        results = self.parallel(action)
        self.assertCountEqual(results, ['saved', 'rejected'])
        self.assertEqual(Attendance.objects.count(), 1)
