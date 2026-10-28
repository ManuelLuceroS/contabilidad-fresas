from io import BytesIO
from decimal import Decimal
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .farm_services import payroll_summary


def payroll_xlsx(payroll):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Resumen'
    summary.append([f'Nómina {payroll.pk} · {payroll.start} a {payroll.end} · {payroll.get_status_display()}'])
    summary.append(['Trabajador', 'Días trabajados', 'Extras (minutos)', 'Pago días (COP)', 'Pago extras (COP)', 'Total (COP)'])
    for w in payroll_summary(payroll):
        summary.append([w['name'], w['days'], w['minutes'], w['base'], w['extra'], w['total']])
    summary.append(['TOTAL GENERAL', None, None, None, None, payroll.total])
    detail = workbook.create_sheet('Detalle diario')
    detail.append(['Trabajador', 'Identificación', 'Fecha', 'Trabajó', 'Extras (minutos)', 'Tarifa día', 'Tarifa hora extra', 'Pago días', 'Pago extras', 'Total', 'Observaciones', 'Regla histórica'])
    for day in payroll.days.order_by('pk'):
        d = day.snapshot
        detail.append([d['worker'], d['identification'], d['date'], 'Sí' if d['worked'] else 'No', d['minutes'],
                       *[Decimal(d[k]) for k in ['daily_rate', 'hourly_rate', 'base', 'extra', 'total']],
                       d['notes'], str(d['rule'])])
    for sheet, header in [(summary, 2), (detail, 1)]:
        sheet.freeze_panes = f'B{header + 1}'
        sheet.auto_filter.ref = f'A{header}:{get_column_letter(sheet.max_column)}{sheet.max_row - (1 if sheet == summary else 0)}'
        for row in sheet:
            for cell in row:
                # Los datos del usuario siempre son texto, incluso si empiezan por '='.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row == header:
                    cell.fill = PatternFill('solid', fgColor='14301F')
                    cell.font = Font(color='FFFFFF', bold=True)
                elif isinstance(cell.value, (Decimal, float)):
                    cell.number_format = '#,##0.00'
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 30 if col in (1, 11, 12) else 22
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = f'1:{header}'
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="nomina-{payroll.pk}.xlsx"'
    return response
